import openpyxl
import tkinter as tk


def create_new_ws():
    ws.title = "Incomes"
    ws1 = wb.create_sheet("Expenses")
    return ws1


def add_heading():
    heading = ["Date", "Amount", "Type"]
    ws.append(heading)
    ws1.append(heading)


def write_to_ws(date, amount, type):
    new_row = ws.append([date, amount, type])
    return new_row


if __name__ == "__main__":
    wb = openpyxl.Workbook()
    ws = wb.active
    ws1 = create_new_ws()
    add_heading()
    date = input("Enter date: ")
    amount = input("Enter amount: ")
    type = input("Enter type: ")
    new_row = write_to_ws(date, amount, type)
    wb.save("quyet.xlsx")
